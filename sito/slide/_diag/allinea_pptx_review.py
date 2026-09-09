# _diag/allinea_pptx_review.py
# Allinea il deck PowerPoint alle correzioni fatte sul beamer nella revisione
# del 2026-09-08/09 (vedi backups/perceiver_presentation_eng_BEFORE_review_20260908.tex
# per il "prima"). Legge il backup BEFORE_review_20260909 e riscrive il deck:
# rieseguibile, ma perde le modifiche fatte a mano dopo il backup.
#
#   python _diag/allinea_pptx_review.py
#
# Slide toccate: 13 (tabella), 34 (figura + testi), 37, 38, 40, 41 (testi).
# Le formule (slide 6/9/26/27) NON si toccano: sono gia' allineate o fatte a mano.

import os
import sys

from pptx import Presentation
from pptx.util import Inches

HERE = os.path.dirname(os.path.abspath(__file__))
SLIDE_DIR = os.path.normpath(os.path.join(HERE, ".."))
SRC = os.path.join(SLIDE_DIR, "backups", "Perceiver & Perceiver IO_BEFORE_review_20260909.pptx")
DST = os.path.join(SLIDE_DIR, "Perceiver & Perceiver IO.pptx")
FIG = os.path.join(SLIDE_DIR, "figure", "attention_maps_v2.png")

BULLET = "•  "


def set_paragraph_text(p, text):
    """Mette `text` nel primo run del paragrafo (ne eredita il formato) e toglie gli altri."""
    runs = p.runs
    if not runs:
        p.add_run().text = text
        return
    runs[0].text = text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)


def replace_in_frame(tf, old, new, where):
    """Sostituisce `old` con `new` nel paragrafo che lo contiene. Esattamente uno."""
    hits = [p for p in tf.paragraphs if old in "".join(r.text for r in p.runs)]
    if len(hits) != 1:
        raise SystemExit(f"{where}: trovate {len(hits)} occorrenze di {old!r} (attesa 1)")
    p = hits[0]
    set_paragraph_text(p, "".join(r.text for r in p.runs).replace(old, new))


def shape_by_id(slide, shape_id):
    for sh in slide.shapes:
        if sh.shape_id == shape_id:
            return sh
    raise SystemExit(f"shape id {shape_id} non trovata sulla slide")


def set_bullets(tf, lines):
    """Riscrive i paragrafi di una casella bullet: tanti paragrafi quante righe."""
    paras = tf.paragraphs
    if len(paras) < len(lines):
        raise SystemExit("piu' righe che paragrafi disponibili")
    for p, line in zip(paras, lines):
        set_paragraph_text(p, BULLET + line)
    for p in paras[len(lines):]:
        p._p.getparent().remove(p._p)


def main():
    if os.path.exists(os.path.join(SLIDE_DIR, "~$Perceiver & Perceiver IO.pptx")):
        raise SystemExit("Il deck e' aperto in PowerPoint: chiudilo.")
    prs = Presentation(SRC)
    S = lambda n: prs.slides[n - 1]

    # --- 13: Implementation (tabella) ------------------------------------------
    tbl = shape_by_id(S(13), 6).table
    replace_in_frame(tbl.cell(4, 2).text_frame, "batches 32–128", "batches 32–64", "s13 batch")
    set_paragraph_text(tbl.cell(5, 2).text_frame.paragraphs[0],
                       "Validation split carved from train, fixed seed, attention-map extraction, declarative run registry")

    # --- 34: Attention maps (figura v2 + testi) ----------------------------------
    s = S(34)
    shape_by_id(s, 4).text_frame.paragraphs[0].runs[0].text = "Attention Maps: Where the Latents Look"
    for r in shape_by_id(s, 4).text_frame.paragraphs[0].runs[1:]:
        r._r.getparent().remove(r._r)
    for pid in (19, 20):                      # le due figure v1
        el = shape_by_id(s, pid)._element
        el.getparent().remove(el)
    # area libera fra la banda (finisce a 2.61in) e le intestazioni (5.55in)
    top, height = Inches(2.72), Inches(2.72)
    pic = s.shapes.add_picture(FIG, Inches(0), top, height=height)
    band_left, band_w = Inches(0.85), Inches(12.01)
    pic.left = int(band_left + (band_w - pic.width) / 2)
    set_paragraph_text(shape_by_id(s, 22).text_frame.paragraphs[0],
                       "With coordinates the latents pick locations; without, they can only pick colours.")
    set_paragraph_text(shape_by_id(s, 23).text_frame.paragraphs[0], "Fourier PE (e01, 71.63%)")
    set_bullets(shape_by_id(s, 24).text_frame, [
        "Each latent attends to a few pixels at specific locations",
        "Mean entropy 5.84 bits over 1,024 pixels (uniform: 10 bits)",
    ])
    set_paragraph_text(shape_by_id(s, 25).text_frame.paragraphs[0], "No PE (e29, 32.36%)")
    set_bullets(shape_by_id(s, 26).text_frame, [
        "Same-colour pixels are indistinguishable: attention follows colour boundaries",
        "9.83 bits, 98% of uniform — spread over the whole image",
    ])

    # --- 37: Hardware -------------------------------------------------------------
    replace_in_frame(shape_by_id(S(37), 10).text_frame,
                     "about eight points below the paper", "about ten points below the paper", "s37")

    # --- 38: Possible improvements -----------------------------------------------
    replace_in_frame(shape_by_id(S(38), 10).text_frame,
                     "Mixed precision, FP16 or BF16", "More seeds per ablation, to narrow the 2.78-point band", "s38 bf16")
    replace_in_frame(shape_by_id(S(38), 12).text_frame,
                     "Subword tokenization for language", "Longer byte-level pre-training: more epochs, larger corpus", "s38 subword")

    # --- 40: Conclusions ----------------------------------------------------------
    replace_in_frame(shape_by_id(S(40), 10).text_frame,
                     "GLUE sits about eight points below the paper, and pre-training scale does not close it",
                     "GLUE sits about ten points below the paper, at a tenth of the parameters and pre-training", "s40")

    # --- 41: Thank you ------------------------------------------------------------
    replace_in_frame(shape_by_id(S(41), 6).text_frame,
                     "Code and models available upon request", "Code and reproduction instructions shared with the course", "s41")

    prs.save(DST)
    print("salvato:", DST)


if __name__ == "__main__":
    main()
