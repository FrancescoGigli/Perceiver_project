# _diag/allinea_rev_20260927.py
# Allinea il PPTX alla revisione del 27/09/2026, come il beamer, la dispensa e la
# lezione: tutto quello che il deck dice sul progetto deve tornare con il codice e i log.
#   - testi e celle: il seed sposta anche lo split di validation; e02 e' atteso per
#     costruzione (la PE si concatena prima di permutare); e16 scondivide solo i blocchi
#     latenti; CoLA in accuratezza e' la classe maggioritaria, in Matthews la media e' 62.49;
#     la ResNet-18 ha la sua ricetta SGD; e23, e24, e27, e28 sono training instabili;
#     ModelNet40 ha un'altra configurazione e nessuna banda; il paper confronta LAMB con
#     SGD, non con Adam; batch e TPU dei paper; ResNet-50 a 77.6 nella Tab. 1.
#   - formula della slide 12: trust ratio ||w|| / ||u + lambda w||.
#   - otto figure rigenerate da _diag/figure_risultati42.py.
#
#   python _diag/allinea_rev_20260927.py
#
# Legge SEMPRE da backups/..._BEFORE_rev_20260927.pptx e riscrive il deck: e'
# rieseguibile, ma una modifica fatta a mano in PowerPoint dopo l'ultima esecuzione
# andrebbe persa. Ogni testo viene sostituito solo se quello attuale e' quello atteso:
# altrimenti lo script si ferma senza salvare. L'immagine mc:Fallback della formula
# resta quella vecchia finche' PowerPoint non risalva il file.

import os
import sys

from pptx import Presentation

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from aggiorna_pptx_42 import swap_pic  # noqa: E402

SLIDE_DIR = os.path.normpath(os.path.join(HERE, ".."))
DECK = os.path.join(SLIDE_DIR, "Perceiver & Perceiver IO.pptx")
SORGENTE = os.path.join(SLIDE_DIR, "backups", "Perceiver & Perceiver IO_BEFORE_rev_20260927.pptx")

NS_MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
NS_P = "http://schemas.openxmlformats.org/presentationml/2006/main"
NS_M = "http://schemas.openxmlformats.org/officeDocument/2006/math"

# (slide, id della shape, paragrafo): (testo atteso, testo nuovo)
TESTI = {(5, 68, 0): ('N(N) + O(M)', 'O(N) + O(M)'),
 (11, 12, 0): ('A Transformer pays for every pair of input tokens. The Perceiver pays once to read the '
               'input, then works on a fixed number of latents.',
               'A Transformer pays for every pair of input tokens. The Perceiver pays M·N per input read, '
               'then works on a fixed number of latents.'),
 (11, 15, 0): ('•  The whole input is read once, through cross-attention',
               '•  Only cross-attention reads the input: O(MN) per read'),
 (11, 15, 1): ('•  Every repeat after that stays inside the N latent slots',
               '•  Every latent block stays inside the N latent slots'),
 (11, 15, 3): ('•  That is why 1,024 pixels and 512 bytes end up costing the same',
               '•  Only the reads grow with M; the latent blocks do not'),
 (12, 23, 0): ('•  The cross-attention bottleneck produces high-variance gradients and plain Adam '
               'destabilises with large batches, so the paper uses LAMB',
               '•  The paper found LAMB easier to optimise than SGD, the usual choice for CNNs on ImageNet; '
               'it does not test Adam'),
 (12, 23, 1): ('•  At batch 64 that advantage disappears — and LAMB is exactly what makes the latent init '
               'scale so dangerous: at 1.0 the run collapses to 52.08%',
               '•  Latent init scale 1.0 collapses our run to 52.08%; in the paper (LAMB too) it costs about '
               '1 point: cause not isolated'),
 (13, 8, 0): ("The mechanism is the paper's, unchanged. What differs is scale, data budget and hardware.",
              "The mechanism is the paper's; scale, budget and a few choices (MLP 4×) differ."),
 (14, 9, 0): ('Not from the budget: ModelNet40 scale and shift move the whole cloud, and the normalisation '
              'undoes them (the paper jitters each point); the MLM masks single bytes, the paper whole '
              'words.',
              'Not from the budget: MLP 4× (paper 1×); ModelNet40 scale and shift undone by normalisation '
              '(the paper jitters each point); single-byte masking, the paper whole words.'),
 (16, 16, 0): ('Twenty-four runs on CIFAR-10. The best reaches 72.91%, the baseline 71.63% — and most of the '
               'gaps between them are smaller than the noise.',
               'Twenty-four runs on CIFAR-10. The best reaches 72.91%, the baseline 71.63% — a 1.28-point '
               'gap, inside the noise.'),
 (16, 18, 1): ('•  Latent init scale at 1.0: −19.55', '•  Latent init scale at 1.0: −19.55, an unstable run'),
 (16, 18, 2): ('•  Sixteen Fourier bands instead of 64: −9.05', '•  Sixteen bands: −9.05, a collapsed run'),
 (18, 18, 2): ('•  Which is convenient: sharing is also what keeps memory affordable on a single GPU',
               '•  Sharing saves weights, yet e16 also ran on the same GPU at batch 64'),
 (19, 14, 2): ('•  The learned-PE runs use a single cross-attend: a learned table is tied to one input '
               'layout',
               '•  The learned-PE runs use one cross-attend, as in the paper: with 8 it was unstable'),
 (20, 13, 0): ('Three runs, identical in everything but the seed, land 2.78 points apart. Read every other '
               'comparison against that.',
               'Three runs that differ in seed and validation split land 2.78 points apart. Read every other '
               'comparison against that.'),
 (20, 15, 1): ('•  Nothing else changed: same data, same schedule, same 120 epochs',
               '•  The seed also moves the validation split; test set unchanged'),
 (20, 17, 0): ('•  Any gap smaller than 2.78 points is variance, not an effect',
               '•  A gap below 2.78 points cannot be told apart from variance'),
 (20, 17, 1): ('•  Twelve of the 24 CIFAR-10 runs clear the band, ten measured from the three-seed mean '
               '(70.48%); the rest are reported as inconclusive',
               '•  Twelve of the 21 CIFAR-10 ablations clear the band, ten measured from the three-seed mean '
               '(70.48%); the rest are reported as inconclusive'),
 (20, 17, 2): ('•  It cost two runs, and it changed the status of all the others',
               '•  It cost two runs and changed the status of the CIFAR-10 runs'),
 (21, 8, 0): ('The same network, fed 2,048 points in 3D instead of pixels. Only the positional encoding '
              'changes dimension.',
              'The same architecture, fed 2,048 points in 3D instead of pixels, with its own configuration '
              '(table).'),
 (21, 12, 0): ("•  87.36% against the paper's 85.7% — above it with a batch sixteen times smaller, and still "
               'above at the last epoch (86.43%)',
               "•  87.36% against the paper's 85.7%, epoch picked on test as in the paper; 86.43% at the "
               'last epoch'),
 (21, 12, 1): ('•  ModelNet40 is small and its objects arrive canonically aligned, so the large-batch '
               'advantage that dominates ImageNet counts for little here',
               '•  Different configuration (6 bands, batch 32, stepped lr) and no noise band: the point is '
               'generality'),
 (22, 9, 1): ('•  Plus translation: 87.20%, the same data as the first run: 0.16 is seed noise',
              '•  Plus translation: 87.20%, the same data as the first run: 0.16 is noise'),
 (22, 11, 0): ('•  ModelNet40 objects arrive canonically aligned, all facing the same way',
               '•  ModelNet40 objects are probably stored upright (not verified)'),
 (23, 17, 0): ('Neither Perceiver paper benchmarks CIFAR-10, so we trained the reference ourselves: a '
               'ResNet-18, same split, same epochs, same budget.',
               'Neither Perceiver paper benchmarks CIFAR-10, so we trained the reference ourselves: a '
               'ResNet-18, same split and epochs, own SGD recipe.'),
 (23, 21, 0): ('•  On ImageNet the paper beats ResNet-50 by 4.5 points: the scale is the variable, not the '
               'architecture',
               '•  On ImageNet the paper edges ResNet-50 by 0.4 (4.5 over ResNet-50 with Fourier features)'),
 (23, 21, 1): ('•  It never claimed accuracy — it claimed generality',
               '•  It claimed generality at comparable accuracy'),
 (23, 21, 2): ('•  Here that prior costs 22 points; the paper says it stops paying at ImageNet scale',
               '•  Here prior plus recipe cost 22 points; at ImageNet scale the prior stops paying'),
 (27, 13, 0): ('On a single label out of ten, the query decoder and mean pooling are the same thing — and '
               'the measurement says so.',
               'On a single label out of ten, our 2.78-point band cannot separate the query decoder from '
               'mean pooling.'),
 (27, 17, 0): ('•  71.79% against 71.63%: 0.16 points, a fraction of the 2.78-point seed spread',
               '•  71.79% vs 71.63% (+0.16, in band), but the recipe changed too: cosine at 1e-3'),
 (30, 18, 0): ("The mean over the eight tasks is 71.13 against the paper's 81.0 — and two of the eight "
               'numbers are not learning at all.',
               "The mean is 71.13 against the paper's 81.0; with CoLA in Matthews, as in the paper, it is "
               '62.49.'),
 (30, 22, 2): ('•  STS-B is reported as Pearson × 100, the official GLUE metric',
               '•  STS-B is reported as Pearson × 100, as in the paper'),
 (31, 19, 1): ('•  The two are chosen at the extremes of data size — SST-2 has 67,000 examples, RTE has '
               '2,500',
               '•  One large and one small task — SST-2 has 67,000 examples, RTE has 2,500'),
 (31, 20, 0): ('And the prediction holds', 'What the controls show'),
 (31, 21, 2): ('•  Transfer is worth five times more where the data abounds: 2,500 examples cannot exploit '
               'the encoder they were handed',
               '•  RTE from scratch is the majority class (52.71%): not learned well either way; one run '
               'each'),
 (32, 18, 0): ('What replicates', 'What agrees'),
 (32, 19, 1): ('•  The gain sits on the small tasks (STS-B +16.9, RTE +7.6, MRPC +6.1); the large ones lose '
               '2 to 3.6',
               '•  The gain sits on the small tasks (STS-B +16.9, RTE +7.6, MRPC +6.1); SST-2, QQP, MNLI '
               'lose 1.9–3.6'),
 (32, 19, 2): ('•  Sign and order of magnitude hold; the absolute level stays about eight points below',
               '•  Same sign, one run per task, no noise band; with CoLA in Matthews, 16 points below'),
 (32, 20, 0): ('Why the win is understated', 'Caveats, both ways'),
 (32, 21, 1): ('•  It wins anyway, with CoLA at the majority class in both models',
               '•  Ahead on one run per task; CoLA at the majority class in both'),
 (32, 21, 2): ('•  And a ninth task would cost one query, not another model',
               '•  A ninth task costs one query and one linear head'),
 (33, 18, 0): ('There is no early stopping: every run trains the full schedule and reports the checkpoint '
               'that was best on validation.',
               'No early stopping: every run trains its full schedule and reports its best epoch '
               '(validation; test on ModelNet40).'),
 (33, 20, 0): ('•  Fine-tuning peaks within a handful of epochs: SST-2 at 7, the MLM at 10',
               '•  SST-2 fine-tuning peaks at epoch 7 of 10; the MLM at its last, 10 of 10'),
 (34, 24, 0): ('•  Each latent attends to a few pixels at specific locations',
               '•  One image: each latent picks a few pixels; latents overlap'),
 (35, 8, 0): ('Forty-two runs, all completed. Every number below traces to a results.json — none is quoted '
              'from memory.',
              'Forty-two runs, all completed. Every number below traces to a log, a results.json or the '
              'papers.'),
 (35, 10, 2): ('•  One multitask model beats eight separate ones, 74.05 against 71.13',
               '•  Multitask 74.05 vs 71.13 separate: one run each, no noise band'),
 (36, 13, 0): ('On the modality with the weakest domain prior we beat the paper. On the one with the '
               'strongest, a convolutional net of the same size beats us.',
               "Point clouds: above the paper's number. Images: a CNN with its own recipe is 22 points "
               'ahead.'),
 (36, 15, 1): ('•  With a batch 16× smaller: the dataset is small and canonically aligned',
               '•  Epoch picked on test, as in the paper; other config, no noise band'),
 (36, 17, 0): ('•  93.61% for a ResNet-18 on the same budget, 71.63% for the Perceiver: 21.98 points below',
               '•  93.61% for a ResNet-18 with its own SGD recipe, 71.63% for the Perceiver: 21.98 below'),
 (36, 18, 0): ('On ImageNet the paper beats ResNet-50 by 4.5 points: the prior stops paying at scale, and '
               'that relation is the claim being replicated.',
               'On ImageNet the paper edges ResNet-50 by 0.4 points (4.5 over its Fourier-feature version); '
               'data scale is not varied here.'),
 (37, 10, 0): ('•  LAMB was built for large batches; at 64 it is sensitive to the latent init scale',
               '•  Init scale 1.0 collapses our run; in the paper (LAMB too) it costs ~1 point'),
 (37, 10, 2): ('•  Half or more of the CIFAR-10 ablations land inside the noise band',
               '•  9 of 21 CIFAR-10 ablations sit inside the band (11 vs seed mean)'),
 (37, 10, 3): ('•  The GLUE level stays about ten points below the paper',
               "•  GLUE, CoLA in Matthews: 62.49 against the paper's 81.0"),
 (37, 12, 0): ('•  Weight sharing to cut memory — and it costs nothing measurable',
               '•  Weight sharing: unsharing (3× weights) gained nothing measurable'),
 (38, 10, 1): ('•  More seeds per ablation, to narrow the 2.78-point band',
               '•  More seeds per ablation, for a better noise estimate'),
 (39, 6, 0): ('Everything measured, split by whether it survived the noise band.',
              'Everything measured; only CIFAR-10 has a noise band.'),
 (39, 8, 4): ('•  One multitask model beats eight separate ones: 74.05 against 71.13',
              '•  Multitask 74.05 vs 71.13: same sign as the paper, no noise band'),
 (39, 10, 0): ('•  A ResNet-18 on the same budget reaches 93.61% against 71.63%',
               '•  A ResNet-18, own SGD recipe, reaches 93.61% against 71.63%'),
 (39, 10, 2): ('•  Two GLUE numbers are the majority class, not learning',
               '•  CoLA is the majority class; MRPC beats it by 9 of 408'),
 (39, 10, 4): ('•  Twelve of the 24 CIFAR-10 runs are indistinguishable from noise, fourteen against the '
               'seed mean',
               '•  Nine of the 21 CIFAR-10 ablations are indistinguishable from noise, eleven against the '
               'seed mean'),
 (40, 8, 1): ('•  Complexity linear in the input size',
              '•  Complexity linear (by construction, not measured)'),
 (40, 8, 2): ('•  Weight sharing: three times fewer parameters at no measurable cost',
              '•  Weight sharing: unsharing (3× weights) gained nothing measurable'),
 (40, 8, 4): ('•  Output queries: one more task costs one query, not another model',
              '•  Output queries: a new task costs one query and one linear head'),
 (40, 10, 0): ('•  At this scale the domain prior is worth 21.98 points',
               '•  At this scale prior plus recipe are worth 21.98 points'),
 (40, 10, 2): ('•  GLUE sits about ten points below the paper, at a tenth of the parameters and pre-training',
               '•  GLUE: 62.49 with CoLA in Matthews vs 81.0, at a tenth of the parameters'),
 (40, 10, 3): ('•  LAMB at batch 64 is sensitive to the latent init scale',
               '•  Init scale 1.0 collapses training (paper: ~1 point)'),
 (40, 11, 0): ("The Perceiver's advantage is not accuracy — it is having no prior on the domain. At CIFAR-10 "
               "scale that prior is worth 22 points; the paper's own claim is that it stops paying at "
               'ImageNet scale.',
               "The Perceiver's advantage is not accuracy but having no domain prior. At CIFAR-10 scale "
               'prior plus recipe are worth 22 points; on ImageNet the paper shows the prior stops paying.')}

# (slide, id della tabella, riga, colonna): (testo atteso, testo nuovo)
CELLE = {(4, 407, 1, 1): ('256', '196'),
 (13, 6, 2, 1): ('Large vision/language: up to 512×1024 latents, 201M-param IO',
                 'Large runs: latents up to 2048×512 (IO flow), 201M-param IO'),
 (13, 6, 4, 1): ('Distributed accelerators, batches up to 512–1024',
                 'Distributed accelerators, batches 512–8192'),
 (17, 6, 6, 5): ('real effect', 'collapsed'),
 (17, 6, 7, 5): ('real effect', 'collapsed'),
 (35, 6, 1, 2): ('72.91%', '72.91% (in band)'),
 (35, 6, 5, 2): ('71.13 mean', '71.13 / 62.49 MCC'),
 (36, 11, 1, 0): ('ResNet-18, same budget', 'ResNet-18, own recipe'),
 (37, 6, 0, 1): ('Original paper (Google)', 'Papers (DeepMind)'),
 (37, 6, 1, 1): ('TPU pods (up to 64 TPUs)', 'TPUs (Perceiver: not stated; IO: 64–256)'),
 (37, 6, 2, 1): ('up to 1024', '512–8192'),
 (37, 6, 3, 1): ('256–512', '512 (ImageNet), 2048 (IO flow)')}

# (slide, id della shape con la formula): coppie (m:t atteso, m:t nuovo), in ordine
OMML = {(12, 15): [('𝑟', '𝑟'), ('=', '='), ('𝑤', '𝑤'), ('𝑤', 'û'), ('+', '+'), ('𝜂', '𝜆'), ('·û', '𝑤')]}

# (slide, id della figura): nuova figura in figure/, rigenerata da figure_risultati42.py
FIGURE = {
    (16, 14): "cifar10_ablation_correct.png",
    (20, 11): "noise_band_correct.png",
    (22, 7): "modelnet40_correct.png",
    (23, 15): "cnn_baseline_correct.png",
    (27, 11): "perceiver_vs_io_correct.png",
    (30, 16): "glue_8task_correct.png",
    (31, 15): "pretraining_value_correct.png",
    (32, 15): "multitask_glue_correct.png",
}


def cerca(shapes, sid):
    for sh in shapes:
        if sh.shape_id == sid:
            return sh
        if sh.shape_type == 6:  # gruppo
            trovata = cerca(sh.shapes, sid)
            if trovata is not None:
                return trovata
    return None


def shape(slide, sid):
    sh = cerca(slide.shapes, sid)
    if sh is None:
        raise SystemExit(f"shape {sid} non trovata")
    return sh


def testo(p):
    return "".join(r.text for r in p.runs)


def sostituisci(p, nuovo, dove):
    """Testo del paragrafo = primo run (ne eredita il formato); gli altri run spariscono."""
    if not p.runs:
        raise SystemExit(f"{dove}: paragrafo senza run")
    p.runs[0].text = nuovo
    for r in p.runs[1:]:
        r._r.getparent().remove(r._r)


def mt_formula(slide, sid):
    """Gli m:t della formula OMML della shape `sid`, dentro il ramo mc:Choice."""
    candidati = []
    for choice in slide.element.iter("{%s}Choice" % NS_MC):
        for c in choice.iter("{%s}cNvPr" % NS_P):
            if c.get("id") == str(sid):
                sp = c.getparent().getparent()
                if next(sp.iter("{%s}oMath" % NS_M), None) is not None:
                    candidati.append(sp)
    if len(candidati) != 1:
        raise SystemExit(f"formula {sid}: trovate {len(candidati)} shape OMML")
    return list(candidati[0].iter("{%s}t" % NS_M))


def main():
    if os.path.exists(os.path.join(SLIDE_DIR, "~$Perceiver & Perceiver IO.pptx")):
        raise SystemExit("Il deck e' aperto in PowerPoint: chiudilo.")
    prs = Presentation(SORGENTE)
    diversi, lavori = [], []
    for (n, sid, i), (vecchio, nuovo) in TESTI.items():
        p = shape(prs.slides[n - 1], sid).text_frame.paragraphs[i]
        lavori.append((p, vecchio, nuovo, f"slide {n} [{sid}] p{i}"))
    for (n, sid, r, c), (vecchio, nuovo) in CELLE.items():
        p = shape(prs.slides[n - 1], sid).table.cell(r, c).text_frame.paragraphs[0]
        lavori.append((p, vecchio, nuovo, f"slide {n} tabella [{sid}] ({r},{c})"))
    for p, vecchio, _, dove in lavori:
        if testo(p) != vecchio:
            diversi.append(f"{dove}: trovato {testo(p)!r}")
    formule = []
    for (n, sid), coppie in OMML.items():
        mt = mt_formula(prs.slides[n - 1], sid)
        if [t.text for t in mt] != [v for v, _ in coppie]:
            diversi.append(f"slide {n} formula [{sid}]: trovato {[t.text for t in mt]!r}")
        formule.append((mt, coppie))
    for (n, sid) in FIGURE:
        sh = shape(prs.slides[n - 1], sid)
        if sh.shape_type != 13:
            diversi.append(f"slide {n} [{sid}]: non e' una figura")
    if diversi:
        raise SystemExit("Testo diverso da quello atteso, niente salvato:\n  " + "\n  ".join(diversi))
    for p, _, nuovo, dove in lavori:
        sostituisci(p, nuovo, dove)
    for mt, coppie in formule:
        for t, (_, nuovo) in zip(mt, coppie):
            t.text = nuovo
    for (n, sid), nome in FIGURE.items():
        swap_pic(prs.slides[n - 1], sid, nome)
    prs.save(DECK)
    print(f"{len(lavori)} testi, {len(formule)} formula e {len(FIGURE)} figure aggiornati: {DECK}")


if __name__ == "__main__":
    main()
