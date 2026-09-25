# _diag/correggi_caveat_20260925.py
# Allinea il PPTX alla revisione del codice del 25/09/2026, come il beamer:
#   - ModelNet40 e testo girano con 6 bande Fourier (il paper: 64);
#   - scala e traslazione di ModelNet40 muovono la nuvola intera e la
#     normalizzazione le annulla: mn03 e' di fatto una replica di mn01;
#   - l'MLM si legge contro la tabella dei vicini (42,96%, baseline_mlm.py),
#     non contro il caso uniforme; il mascheramento e' su byte singoli;
#   - la banda si misura da e01, il piu' fortunato dei tre seed (media 70,48%);
#   - il +2,92 del multitask sta tutto nei task piccoli;
#   - ModelNet40: epoca scelta sul test, l'ultima fa 86,43%.
#
#   python _diag/correggi_caveat_20260925.py
#
# Legge SEMPRE da backups/..._BEFORE_caveat_20260925.pptx e riscrive il deck: e'
# rieseguibile, ma una modifica fatta a mano in PowerPoint dopo l'ultima
# esecuzione andrebbe persa. Ogni frase viene sostituita solo se il testo attuale
# e' quello atteso: altrimenti lo script si ferma senza salvare.

import os

from pptx import Presentation

HERE = os.path.dirname(os.path.abspath(__file__))
SLIDE_DIR = os.path.normpath(os.path.join(HERE, ".."))
DECK = os.path.join(SLIDE_DIR, "Perceiver & Perceiver IO.pptx")
SORGENTE = os.path.join(SLIDE_DIR, "backups", "Perceiver & Perceiver IO_BEFORE_caveat_20260925.pptx")

# (slide, id della shape, paragrafo): (testo atteso, testo nuovo)
TESTI = {
    (14, 8, 0): (
        "Every difference from the paper follows from one GPU instead of 512 TPU cores. "
        "The architecture itself is untouched.",
        "Most differences from the paper follow from one GPU instead of 512 TPU cores. "
        "The architecture itself is untouched."),
    (14, 9, 0): (
        "Cross-attention bottleneck, weight-shared latent transformer, then mean pooling or an "
        "output-query decoder. The input is raw pixels, exactly as in the paper.",
        "Not from the budget: ModelNet40 scale and shift move the whole cloud, and the "
        "normalisation undoes them (the paper jitters each point); the MLM masks single bytes, "
        "the paper whole words."),
    (20, 15, 0): (
        "•  Seed 42 gives 71.63%, seed 2 gives 70.97%, seed 1 gives 68.85%",
        "•  Seed 42 gives 71.63%, seed 2 gives 70.97%, seed 1 gives 68.85%: the reference "
        "is the luckiest of the three"),
    (20, 17, 1): (
        "•  Twelve of the 24 CIFAR-10 runs clear the band; the other twelve are reported as "
        "inconclusive instead of dressed up as trends",
        "•  Twelve of the 24 CIFAR-10 runs clear the band, ten measured from the three-seed "
        "mean (70.48%); the rest are reported as inconclusive"),
    (21, 10, 1): (
        "•  The Fourier encoding becomes 3D, over (x, y, z) instead of (u, v)",
        "•  The Fourier encoding becomes 3D over (x, y, z), with 6 bands per axis "
        "(the paper uses 64)"),
    (21, 12, 0): (
        "•  87.36% against the paper's 85.7% — above it, with a batch sixteen times smaller",
        "•  87.36% against the paper's 85.7% — above it with a batch sixteen times smaller, "
        "and still above at the last epoch (86.43%)"),
    (22, 6, 0): (
        "Translation is free. Rotation costs 13.29 points — and the reason is the positional "
        "encoding again.",
        "Rotation costs 13.29 points — the positional encoding again. Scale and translation "
        "never reach the network."),
    (22, 9, 1): (
        "•  Plus translation: 87.20%, a loss of 0.16 points",
        "•  Plus translation: 87.20%, the same data as the first run: 0.16 is seed noise"),
    (22, 11, 2): (
        "•  Translation is harmless because the coordinates are re-centred anyway",
        "•  Scale and translation move the whole cloud, and the re-centring undoes them; "
        "the paper jitters each point"),
    (28, 14, 2): (
        "•  86.68% correct against a chance level of 0.39% — about 222 times better than guessing",
        "•  86.68% correct: 44 points above a lookup of the two neighbouring bytes, 68 above "
        "always answering a space"),
    (32, 19, 1): (
        "•  The absolute level stays about eight points below and does not close",
        "•  The gain sits on the small tasks (STS-B +16.9, RTE +7.6, MRPC +6.1); "
        "the large ones lose 2 to 3.6"),
    (32, 21, 1): (
        "•  It wins anyway",
        "•  It wins anyway, with CoLA at the majority class in both models"),
    (32, 19, 2): (
        "•  But sign and order of magnitude hold, and that relation is what Table 2 claims",
        "•  Sign and order of magnitude hold; the absolute level stays about eight points below"),
    (35, 10, 0): (
        "•  ModelNet40 beats the paper by 1.66 points",
        "•  ModelNet40 above the paper: +1.66 (last epoch +0.73)"),
    (35, 10, 1): (
        "•  The byte-level MLM reaches 86.68% against a 0.39% chance level",
        "•  Byte-level MLM: 86.68%, 44 above a neighbour lookup"),
    (35, 12, 2): (
        "•  53 tests passing, and a 2.78-point noise band measured before any comparison was read",
        "•  55 tests passing, and a 2.78-point noise band measured before any comparison was read"),
    (36, 15, 0): (
        "•  85.70% in the paper, 87.36% here: 1.66 points above it",
        "•  85.70% in the paper, 87.36% here: 1.66 points above, 0.73 at the last epoch"),
    (36, 15, 1): (
        "•  With a batch sixteen times smaller. The dataset is small and canonically aligned, "
        "so the large-batch advantage counts for little",
        "•  With a batch 16× smaller: the dataset is small and canonically aligned"),
    (37, 10, 2): (
        "•  Half the CIFAR-10 ablations land inside the noise band",
        "•  Half or more of the CIFAR-10 ablations land inside the noise band"),
    (39, 8, 2): (
        "•  The byte-level MLM works: 86.68% against a 0.39% chance level",
        "•  The byte-level MLM works: 86.68%, 44 points above a neighbour lookup"),
    (39, 10, 4): (
        "•  Twelve of the 24 CIFAR-10 runs are indistinguishable from noise",
        "•  Twelve of the 24 CIFAR-10 runs are indistinguishable from noise, fourteen "
        "against the seed mean"),
}

# (slide, id della tabella, riga, colonna): (testo atteso, testo nuovo)
CELLE = {
    (14, 6, 0, 1): ("Original paper (ImageNet)", "Original paper"),
    (14, 6, 4, 1): ("2D Fourier, K = 64 bands/axis (258 dims)",
                    "Fourier, K = 64 bands/axis (ImageNet: 258 dims; ModelNet40 too)"),
    (14, 6, 4, 2): ("2D Fourier, K = 64 bands/axis (258 dims) — identical",
                    "CIFAR-10 identical (258 dims); ModelNet40 and text: K = 6 "
                    "(42 per point, 270 per byte)"),
    (28, 8, 4, 0): ("Mask probability", "Masking"),
    (28, 8, 4, 1): ("15%", "15% of single bytes (paper: words)"),
    (28, 8, 9, 0): ("Chance level", "Lookup of the two neighbours"),
    (28, 8, 9, 1): ("0.39%   (1 / 256)", "42.96%   (no network)"),
    (35, 6, 4, 3): ("0.39% chance", "42.96% lookup"),
}


def shape(slide, sid):
    for sh in slide.shapes:
        if sh.shape_id == sid:
            return sh
    raise SystemExit(f"shape {sid} non trovata")


def testo(p):
    return "".join(r.text for r in p.runs)


def sostituisci(p, nuovo, dove):
    """Testo del paragrafo = primo run (ne eredita il formato); gli altri run spariscono."""
    if not p.runs:
        raise SystemExit(f"{dove}: paragrafo senza run")
    p.runs[0].text = nuovo
    for r in p.runs[1:]:
        r._r.getparent().remove(r._r)


def main():
    if os.path.exists(os.path.join(SLIDE_DIR, "~$Perceiver & Perceiver IO.pptx")):
        raise SystemExit("Il deck e' aperto in PowerPoint: chiudilo.")
    prs = Presentation(SORGENTE)
    diversi = []
    lavori = []
    for (n, sid, i), (vecchio, nuovo) in TESTI.items():
        p = shape(prs.slides[n - 1], sid).text_frame.paragraphs[i]
        lavori.append((p, vecchio, nuovo, f"slide {n} [{sid}] p{i}"))
    for (n, sid, r, c), (vecchio, nuovo) in CELLE.items():
        p = shape(prs.slides[n - 1], sid).table.cell(r, c).text_frame.paragraphs[0]
        lavori.append((p, vecchio, nuovo, f"slide {n} tabella [{sid}] ({r},{c})"))
    for p, vecchio, _, dove in lavori:
        if testo(p) != vecchio:
            diversi.append(f"{dove}: trovato {testo(p)!r}")
    if diversi:
        raise SystemExit("Testo diverso da quello atteso, niente salvato:\n  " + "\n  ".join(diversi))
    for p, _, nuovo, dove in lavori:
        sostituisci(p, nuovo, dove)
    prs.save(DECK)
    print(f"{len(lavori)} testi aggiornati: {DECK}")


if __name__ == "__main__":
    main()
