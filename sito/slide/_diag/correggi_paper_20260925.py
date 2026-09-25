# _diag/correggi_paper_20260925.py
# Allinea il PPTX al confronto con i PDF dei paper (25/09/2026), come il beamer:
#   - solo la prima cross-attention ha pesi propri: i blocchi latenti sono
#     condivisi da tutte le letture, anche dalla prima (didascalia Tab. 7);
#   - l'MLP del paper ha fattore 1 (hidden = D, pp. 6 e 17), il nostro 4;
#   - nelle run su T la profondita' cresce con T (4 blocchi per lettura),
#     mentre nella Tab. 6 del paper i blocchi latenti restano 48;
#   - "512 TPU cores" non compare nei paper: il Perceiver cita al massimo 64 TPU.
#
#   python _diag/correggi_paper_20260925.py
#
# Legge SEMPRE da backups/..._BEFORE_paper_20260925.pptx e riscrive il deck: e'
# rieseguibile, ma una modifica fatta a mano in PowerPoint dopo l'ultima
# esecuzione andrebbe persa. Ogni frase viene sostituita solo se il testo attuale
# e' quello atteso: altrimenti lo script si ferma senza salvare.

import os

from pptx import Presentation

HERE = os.path.dirname(os.path.abspath(__file__))
SLIDE_DIR = os.path.normpath(os.path.join(HERE, ".."))
DECK = os.path.join(SLIDE_DIR, "Perceiver & Perceiver IO.pptx")
SORGENTE = os.path.join(SLIDE_DIR, "backups", "Perceiver & Perceiver IO_BEFORE_paper_20260925.pptx")

# (slide, id della shape, paragrafo): (testo atteso, testo nuovo)
TESTI = {
    (10, 15, 2): (
        "•  Weight sharing reuses the same processor at every iteration — except the first block, "
        "which keeps its own",
        "•  Weight sharing reuses the same processor at every iteration; only the first "
        "cross-attention keeps its own weights"),
    (14, 8, 0): (
        "Most differences from the paper follow from one GPU instead of 512 TPU cores. "
        "The architecture itself is untouched.",
        "Most differences from the paper follow from one GPU instead of a TPU pod. "
        "The mechanism itself is untouched."),
    (16, 20, 0): (
        "•  One cross-attend instead of four gains 1.28 points, which is inside the 2.78-point band",
        "•  One cross-attend instead of four gains 1.28 points, inside the 2.78-point band, "
        "with 4 latent blocks instead of 16: in our code depth grows with T"),
    (37, 8, 0): (
        "Every limitation below is one GPU instead of 512 TPU cores. The architecture is the "
        "paper's; the budget is not.",
        "Every limitation below is one GPU instead of a TPU pod. The mechanism is the "
        "paper's; the budget is not."),
}

# (slide, id della tabella, riga, colonna): (testo atteso, testo nuovo)
CELLE = {
    (14, 6, 6, 0): ("Parameters", "Parameters, MLP width"),
    (14, 6, 6, 1): ("30–200M+", "44.9M (ImageNet); MLP 1× (hidden = D)"),
    (14, 6, 6, 2): ("10.2M (CIFAR-10) · 23.3M (ModelNet40) · 18.9M (IO text)",
                    "10.2M (CIFAR-10) · 23.3M (ModelNet40) · 18.9M (IO text); MLP 4×"),
    (37, 6, 1, 1): ("512 TPU v3 cores", "TPU pods (up to 64 TPUs)"),
    (37, 6, 4, 1): ("30–200M+ params", "44.9M (Perceiver), 201M (IO text)"),
}


def shape(slide, sid):
    for sh in slide.shapes:
        if sh.shape_id == sid:
            return sh
    raise SystemExit(f"shape {sid} non trovata")


def testo(p):
    return "".join(r.text for r in p.runs)


def sostituisci(p, nuovo, dove):
    """Testo del paragrafo = primo run (ne eredita il formato); gli altri run spariscono."""
    if not p.runs:
        raise SystemExit(f"{dove}: paragrafo senza run")
    p.runs[0].text = nuovo
    for r in p.runs[1:]:
        r._r.getparent().remove(r._r)


def main():
    if os.path.exists(os.path.join(SLIDE_DIR, "~$Perceiver & Perceiver IO.pptx")):
        raise SystemExit("Il deck e' aperto in PowerPoint: chiudilo.")
    prs = Presentation(SORGENTE)
    diversi = []
    lavori = []
    for (n, sid, i), (vecchio, nuovo) in TESTI.items():
        p = shape(prs.slides[n - 1], sid).text_frame.paragraphs[i]
        lavori.append((p, vecchio, nuovo, f"slide {n} [{sid}] p{i}"))
    for (n, sid, r, c), (vecchio, nuovo) in CELLE.items():
        p = shape(prs.slides[n - 1], sid).table.cell(r, c).text_frame.paragraphs[0]
        lavori.append((p, vecchio, nuovo, f"slide {n} tabella [{sid}] ({r},{c})"))
    for p, vecchio, _, dove in lavori:
        if testo(p) != vecchio:
            diversi.append(f"{dove}: trovato {testo(p)!r}")
    if diversi:
        raise SystemExit("Testo diverso da quello atteso, niente salvato:\n  " + "\n  ".join(diversi))
    for p, _, nuovo, dove in lavori:
        sostituisci(p, nuovo, dove)
    prs.save(DECK)
    print(f"{len(lavori)} testi aggiornati: {DECK}")


if __name__ == "__main__":
    main()
