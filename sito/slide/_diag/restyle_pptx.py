# Porta le slide 10-40 nello stile che l'utente ha usato a mano nelle prime 9.
#
#   python _diag/restyle_pptx.py          # dalla cartella sito/slide/
#
# Grammatica estratta dalle slide 7/8/9, che sono il riferimento:
#   - NESSUN riquadro: niente card azzurre, niente callout crema, niente barre
#     d'accento. Tutto appoggiato sul bianco.
#   - una FRASE in prosa in cima che dice l'idea (banda blu #156082 o testo
#     scuro), non un'etichetta;
#   - il visivo (figura o tabella) al centro;
#   - in basso due colonne: intestazione blu sz16 bold, corpo sz14-15 #1A1A1A
#     con bullet che sono proposizioni intere, non frammenti.
#
# Tre layout, perche' non tutti i visivi hanno la stessa forma:
#   L1  figura larga al centro + due colonne di prosa sotto  (come la slide 8/9)
#   L2  visivo alto a sinistra + prosa a destra              (tabelle lunghe, grafici alti)
#   L3  sola prosa su due o tre colonne                      (slide teoriche e finali)
#
# Parte dal backup post-allineamento-42 e riscrive il deck, quindi e'
# rieseguibile. Le figure e le tabelle esistenti vengono riusate, non ricreate.

import copy
import os

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Emu, Inches, Pt

QUI = os.path.dirname(os.path.abspath(__file__))
SORGENTE = os.path.join(QUI, "..", "backups",
                        "Perceiver & Perceiver IO_BEFORE_restyle.pptx")
DECK = os.path.join(QUI, "..", "Perceiver & Perceiver IO.pptx")

BLU = RGBColor(0x15, 0x60, 0x82)      # intestazioni e banda: colore delle slide 1-9
INK = RGBColor(0x1A, 0x1A, 0x1A)
BIANCO = RGBColor(0xFF, 0xFF, 0xFF)

X0, W = 0.85, 12.01                   # banda utile in larghezza
LEAD_Y, LEAD_H = 2.03, 0.58

L1_VIS_Y, L1_VIS_H = 2.80, 2.10       # figura al centro
L1_COL_Y, L1_BODY_Y, L1_BODY_H = 5.02, 5.40, 1.75

L2_VIS_X, L2_VIS_W = 0.72, 6.90       # visivo a sinistra
L2_TXT_X, L2_TXT_W = 7.80, 5.06

L3_COL_Y, L3_BODY_Y, L3_BODY_H = 2.95, 3.35, 3.60


# ── primitive ───────────────────────────────────────────────────────────────
def shape(slide, sid):
    for sh in slide.shapes:
        if sh.shape_id == sid:
            return sh
    raise KeyError(f"shape {sid} assente")


def ha_formula(sh):
    """Le caselle-formula contengono oMath dentro un ramo a14:m. python-pptx non
    ne enumera i run, quindi l'unico modo di riconoscerle e' guardare l'XML: se
    le cancellassi, la matematica sparirebbe dalla slide."""
    return "a14:m" in sh._element.xml


def pulisci(slide, tieni=()):
    """Toglie card, barre d'accento e vecchie caselle. Restano il footer, il
    titolo, le figure, le tabelle, le formule e quanto elencato in `tieni`."""
    for sh in list(slide.shapes):
        if sh.shape_id in tieni or sh.shape_id == 2:
            continue
        if sh.shape_type == 13 or sh.has_chart or sh.has_table:
            continue
        if sh.has_text_frame and sh == slide.shapes.title:
            continue
        if ha_formula(sh):
            continue
        sh._element.getparent().remove(sh._element)


def casella(slide, x, y, w, h, righe, sz, bold=False, colore=INK,
            algn=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    for i, riga in enumerate(righe):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = algn
        run = para.add_run()
        run.text = riga
        run.font.size = Pt(sz)
        run.font.bold = bold
        run.font.color.rgb = colore
    return box


def lead(slide, testo, banda=True):
    """La frase-guida in cima. Una frase intera, non un'etichetta."""
    if banda:
        from pptx.enum.shapes import MSO_SHAPE
        r = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(X0),
                                   Inches(LEAD_Y), Inches(W), Inches(LEAD_H))
        r.fill.solid()
        r.fill.fore_color.rgb = BLU
        r.line.fill.background()
        r.shadow.inherit = False
        r.text_frame.text = ""
        casella(slide, 1.01, LEAD_Y + 0.07, 11.68, 0.42, [testo], 15,
                bold=True, colore=BIANCO, algn=PP_ALIGN.CENTER)
    else:
        casella(slide, 1.00, LEAD_Y - 0.03, 11.50, 0.64, [testo], 15,
                bold=True, colore=INK)


def colonne(slide, blocchi, y_tit, y_corpo, h_corpo, sz=14.5):
    """Le colonne in basso: intestazione blu, poi frasi intere."""
    n = len(blocchi)
    larghezza = (W - 0.5 * (n - 1)) / n
    for i, (titolo, righe) in enumerate(blocchi):
        x = X0 + i * (larghezza + 0.5)
        casella(slide, x + 0.10, y_tit, larghezza - 0.10, 0.33, [titolo], 16,
                bold=True, colore=BLU)
        casella(slide, x + 0.10, y_corpo, larghezza - 0.10, h_corpo, righe, sz)


def piazza(sh, x, y, w, h=None):
    """Riposiziona una figura o una tabella rispettando le proporzioni."""
    if sh.shape_type == 13 and h is not None:
        scala = min(Inches(w) / sh.width, Inches(h) / sh.height)
        sh.width, sh.height = int(sh.width * scala), int(sh.height * scala)
        sh.left = Inches(x) + (Inches(w) - sh.width) // 2
        sh.top = Inches(y) + (Inches(h) - sh.height) // 2
    else:
        sh.left, sh.top, sh.width = Inches(x), Inches(y), Inches(w)
        if h is not None:
            sh.height = Inches(h)


def L1(slide, testo_lead, visivo_id, blocchi, banda=True, vis_y=L1_VIS_Y,
       vis_h=L1_VIS_H, sz=14):
    pulisci(slide)
    lead(slide, testo_lead, banda)
    piazza(shape(slide, visivo_id), X0, vis_y, W, vis_h)
    colonne(slide, blocchi, L1_COL_Y, L1_BODY_Y, L1_BODY_H, sz=sz)


def L2(slide, testo_lead, visivo_id, blocchi, banda=True, vis_y=2.82,
       vis_h=3.95, vis_w=L2_VIS_W, vis_x=L2_VIS_X, txt_x=L2_TXT_X,
       txt_w=L2_TXT_W, dy=2.10, sz=14.5):
    pulisci(slide)
    lead(slide, testo_lead, banda)
    piazza(shape(slide, visivo_id), vis_x, vis_y, vis_w, vis_h)
    y = vis_y
    for titolo, righe in blocchi:
        casella(slide, txt_x, y, txt_w, 0.33, [titolo], 16, bold=True, colore=BLU)
        casella(slide, txt_x, y + 0.38, txt_w, dy - 0.45, righe, sz)
        y += dy


def L3(slide, testo_lead, blocchi, banda=True, sz=15, y=L3_COL_Y, h=L3_BODY_H):
    pulisci(slide)
    lead(slide, testo_lead, banda)
    colonne(slide, blocchi, y, y + 0.40, h, sz=sz)


def B(*frasi):
    """Bullet nel formato delle slide 1-9: '•' e due spazi."""
    return ["•  " + f for f in frasi]


def chiusura(slide, testo, y=6.55, sz=14.5):
    """Riga conclusiva: nelle slide 1-9 non ci sono bande in fondo, quindi il
    testo di chiusura e' prosa normale, non un banner."""
    return casella(slide, X0 + 0.16, y, W - 0.32, 0.62, [testo], sz, bold=True)


def altcontent(slide):
    """Le caselle-formula costruite con l'equation editor stanno dentro un
    mc:AlternateContent, che python-pptx NON enumera: slide.shapes non le vede e
    pulisci() non le tocca (bene), ma per spostarle serve lxml."""
    return [c for c in slide.shapes._spTree if c.tag.endswith("}AlternateContent")]


def sposta_alt(el, x, y, w=None, h=None):
    """Riposiziona una shape dentro mc:AlternateContent. Va patchato SIA il ramo
    Choice (quello che PowerPoint renderizza) SIA il Fallback, altrimenti i due
    rami si contraddicono."""
    A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
    for off in el.iter(A + "off"):
        off.set("x", str(Inches(x))); off.set("y", str(Inches(y)))
    if w is not None:
        for ext in el.iter(A + "ext"):
            ext.set("cx", str(Inches(w)))
            if h is not None:
                ext.set("cy", str(Inches(h)))


def figura(slide, nome, x, y, w, h):
    percorso = os.path.join(QUI, "..", "figure", nome)
    pic = slide.shapes.add_picture(percorso, Inches(x), Inches(y))
    scala = min(Inches(w) / pic.width, Inches(h) / pic.height)
    pic.width, pic.height = int(pic.width * scala), int(pic.height * scala)
    pic.left = Inches(x) + (Inches(w) - pic.width) // 2
    pic.top = Inches(y) + (Inches(h) - pic.height) // 2
    return pic


def math_scura(sh):
    """Le formule della slide 26 erano bianche perche' stavano su una placca blu.
    Tolta la placca resterebbero invisibili: qui il bianco dei run matematici
    diventa inchiostro. Nelle slide 1-9 le formule stanno sul bianco, senza placca."""
    A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
    for clr in sh._element.iter(A + "srgbClr"):
        if clr.get("val", "").upper() == "FFFFFF":
            clr.set("val", "1A1A1A")


def dietro(sh, altro):
    """Manda `sh` dietro ad `altro` nello z-order: add_shape appende in fondo,
    quindi una banda disegnata dopo coprirebbe il testo che deve stare sopra."""
    altro._element.addprevious(sh._element)


def togli(slide, *ids):
    for sid in ids:
        try:
            sh = shape(slide, sid)
        except KeyError:
            continue
        sh._element.getparent().remove(sh._element)
